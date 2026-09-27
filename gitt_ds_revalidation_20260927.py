"""Revalidate the provisional GITT-based Dsn/Dsp values.

This script is intentionally read-only with respect to the supplied Neware
workbooks.  It checks the exported D sheet, reconstructs the four voltages
used by Neware from the raw record sheets, and audits the square-root-of-time
assumption over several early-pulse windows.
"""

from __future__ import annotations

import bisect
import csv
import math
import warnings
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook


warnings.filterwarnings("ignore", message="Workbook contains no default style")

ROOT = Path(__file__).resolve().parent
SUMMARY_DIR = ROOT.parent.parent / "tmp" / "gitt_ds_audit"
RAW_DIR = Path(
    r"C:\Users\user\OneDrive\01. 적응형 충전 프로토콜\4. Enertech 파우치셀 실험데이터"
    r"\260918 GITT 실험데이터"
)
OUT_DIR = ROOT / "results" / "260927_gitt_ds_revalidation"


@dataclass(frozen=True)
class Electrode:
    name: str
    mass_g: float
    molar_mass_g_mol: float
    molar_volume_cm3_mol: float
    area_cm2: float
    radius_m: float
    provisional_D_m2_s: float


ELECTRODES = {
    "Anode": Electrode(
        "Anode", 0.01040, 72.066, 31.89, 0.7854, 3.7542e-6, 2.1e-14
    ),
    "Cathode": Electrode(
        "Cathode", 0.01725, 97.87, 19.38, 0.7854, 4.2387e-6, 4.4e-14
    ),
}

CASES = [
    (
        "Anode",
        "v1",
        "260920 Ai2020 Anode GITT_Dsn v1 1-5.xlsx",
        "260918 Anode GITT v1 1-5.xlsx",
        None,
    ),
    (
        "Anode",
        "v2",
        "260920 Ai2020 Anode GITT_Dsn v2 1-6.xlsx",
        "260918 Anode GITT v2 1-6.xlsx",
        None,
    ),
    (
        "Anode",
        "v3",
        "260920 Ai2020 Anode GITT_Dsn v3 1-7.xlsx",
        "260918 Anode GITT v3 1-7.xlsx",
        None,
    ),
    (
        "Cathode",
        "v1",
        "260920 Ai2020 Cathode GITT_Dsp v1 2-6.xlsx",
        "260918 Cathode GITT v1 2-6.xlsx",
        "260918 Cathode GITT v1 2-6_1.xlsx",
    ),
    (
        "Cathode",
        "v2",
        "260920 Ai2020 Cathode GITT_Dsp v2 2-7.xlsx",
        "260918 Cathode GITT v2 2-7.xlsx",
        "260918 Cathode GITT v2 2-7_1.xlsx",
    ),
    (
        "Cathode",
        "v3",
        "260920 Ai2020 Cathode GITT_Dsp v3 2-8.xlsx",
        "260918 Cathode GITT v3 2-8.xlsx",
        "260918 Cathode GITT v3 2-8_1.xlsx",
    ),
]

NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PKG_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def geometric_mean(values: list[float] | np.ndarray) -> float:
    array = np.asarray(values, dtype=float)
    array = array[np.isfinite(array) & (array > 0)]
    return float(np.exp(np.mean(np.log(array)))) if len(array) else float("nan")


def record_sheet_path(book_path: Path) -> str:
    with zipfile.ZipFile(book_path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        relation_targets = {
            node.attrib["Id"]: node.attrib["Target"]
            for node in rels.findall(f"{PKG_REL_NS}Relationship")
        }
        sheets = workbook.find(f"{NS}sheets")
        assert sheets is not None
        for sheet in sheets:
            if sheet.attrib.get("name") == "record":
                rel_id = sheet.attrib[f"{REL_NS}id"]
                target = relation_targets[rel_id].lstrip("/")
                return target if target.startswith("xl/") else f"xl/{target}"
    raise RuntimeError(f"record sheet not found: {book_path}")


def read_summary(path: Path, electrode: str, replicate: str) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["D"]
    sheet.reset_dimensions()
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    workbook.close()
    output: list[dict[str, object]] = []
    for sequence, row in enumerate(rows, start=1):
        if row[0] is None:
            continue
        step_type = str(row[2])
        first_branch = step_type in {"5Rest", "7Rest"}
        if electrode == "Anode":
            direction = "Discharge" if first_branch else "Charge"
        else:
            direction = "Charge" if first_branch else "Discharge"
        output.append(
            {
                "electrode": electrode,
                "replicate": replicate,
                "sequence": sequence,
                "data_point": int(row[1]),
                "step_type": step_type,
                "direction": direction,
                "neware_D_cm2_s": float(row[6]),
                "neware_D_m2_s": float(row[6]) * 1e-4,
            }
        )
    return output


def collect_raw_intervals(
    workbook_paths: list[Path], pulses: list[dict[str, object]]
) -> dict[int, dict[str, object]]:
    ordered = sorted(pulses, key=lambda item: int(item["data_point"]))
    starts = [int(item["data_point"]) for item in ordered]
    states: dict[int, dict[str, object]] = {}
    for index, pulse in enumerate(ordered):
        start = starts[index]
        next_start = starts[index + 1] if index + 1 < len(starts) else None
        if next_start is not None and ordered[index + 1]["direction"] != pulse["direction"]:
            next_start = None
        states[start] = {
            "next_start": next_start,
            "V0": None,
            "V1": None,
            "V2": None,
            "V3": None,
            "last_current_point": None,
            "pulse_points": {},
            "tail_points": [],
        }
    for workbook_path in workbook_paths:
        if not workbook_path.exists():
            continue
        sheet_path = record_sheet_path(workbook_path)
        with zipfile.ZipFile(workbook_path) as archive, archive.open(sheet_path) as stream:
            for _, row in ET.iterparse(stream, events=("end",)):
                if row.tag != f"{NS}row":
                    continue
                values: dict[str, float] = {}
                for cell in row.findall(f"{NS}c"):
                    reference = cell.attrib.get("r", "")
                    column = "".join(ch for ch in reference if ch.isalpha())
                    if column not in {"A", "E", "F"}:
                        continue
                    value_node = cell.find(f"{NS}v")
                    if value_node is not None and value_node.text is not None:
                        values[column] = float(value_node.text)
                if "A" in values and "F" in values:
                    data_point = int(values["A"])
                    position = bisect.bisect_right(starts, data_point) - 1
                    if position >= 0:
                        if position > 0:
                            previous_start = starts[position - 1]
                            previous_state = states[previous_start]
                            if previous_state["next_start"] == data_point:
                                previous_state["V3"] = values["F"]
                                previous_state["tail_points"].append(
                                    (data_point, values["F"])
                                )
                        start = starts[position]
                        state = states[start]
                        next_start = state["next_start"]
                        if data_point == start:
                            state["V0"] = values["F"]
                        if data_point == start + 1:
                            state["V1"] = values["F"]
                        if next_start is not None and start < data_point <= int(next_start):
                            current = values.get("E", 0.0)
                            if abs(current) > 0:
                                state["V2"] = values["F"]
                                state["last_current_point"] = data_point
                                elapsed = data_point - (start + 1)
                                if elapsed <= 300:
                                    state["pulse_points"][elapsed] = (current, values["F"])
                            if data_point == int(next_start):
                                state["V3"] = values["F"]
                            if data_point >= int(next_start) - 600:
                                state["tail_points"].append((data_point, values["F"]))
                row.clear()
    return states


def linear_fit_sqrt_time(
    pulse_points: dict[int, tuple[float, float]], start: int, stop: int
) -> tuple[float, float, int]:
    times: list[float] = []
    volts: list[float] = []
    for second in range(start, stop + 1):
        if second in pulse_points:
            times.append(float(second))
            volts.append(float(pulse_points[second][1]))
    if len(times) < 3:
        return float("nan"), float("nan"), len(times)
    x = np.sqrt(np.asarray(times, dtype=float))
    y = np.asarray(volts, dtype=float)
    slope, intercept = np.polyfit(x, y, 1)
    predicted = slope * x + intercept
    residual = float(np.sum((y - predicted) ** 2))
    total = float(np.sum((y - np.mean(y)) ** 2))
    r2 = 1.0 - residual / total if total > 0 else float("nan")
    return float(slope), float(r2), len(times)


def reconstruct_pulse(
    pulse: dict[str, object], raw: dict[int, dict[str, object]], electrode: Electrode
) -> dict[str, object]:
    p = int(pulse["data_point"])
    output = dict(pulse)
    state = raw[p]
    if any(
        state[key] is None
        for key in ("V0", "V1", "V2", "V3", "last_current_point")
    ):
        output.update({"raw_complete": False})
        return output
    v0 = float(state["V0"])
    v1 = float(state["V1"])
    v2 = float(state["V2"])
    v3 = float(state["V3"])
    tau_s = float(int(state["last_current_point"]) - (p + 1))
    delta_es = abs(v0 - v3)
    delta_et = abs(v1 - v2)
    factor = electrode.mass_g * electrode.molar_volume_cm3_mol / (
        electrode.molar_mass_g_mol * electrode.area_cm2
    )
    reconstructed_cm2_s = (
        4.0 / (math.pi * tau_s) * factor**2 * (delta_es / delta_et) ** 2
        if delta_et > 0 and tau_s > 0
        else float("nan")
    )
    pulse_points = state["pulse_points"]
    output.update(
        {
            "raw_complete": True,
            "V0_V": v0,
            "V1_V": v1,
            "V2_V": v2,
            "V3_V": v3,
            "delta_Es_mV": delta_es * 1000.0,
            "delta_Etau_mV": delta_et * 1000.0,
            "pulse_duration_s": tau_s,
            "reconstructed_D_m2_s": reconstructed_cm2_s * 1e-4,
            "neware_reconstruction_rel_error_pct": 100.0
            * (reconstructed_cm2_s - float(pulse["neware_D_cm2_s"]))
            / float(pulse["neware_D_cm2_s"]),
        }
    )
    slopes: list[float] = []
    for label, start, stop in (
        ("10_60", 10, 60),
        ("30_120", 30, 120),
        ("30_300", 30, 300),
    ):
        slope, r2, count = linear_fit_sqrt_time(pulse_points, start, stop)
        output[f"slope_{label}_V_sqrt_s"] = slope
        output[f"R2_{label}"] = r2
        output[f"n_{label}"] = count
        slopes.append(abs(slope))
        slope_D_cm2_s = (
            4.0
            / (math.pi * tau_s)
            * factor**2
            * (delta_es / (abs(slope) * math.sqrt(tau_s))) ** 2
            if slope and np.isfinite(slope) and tau_s > 0
            else float("nan")
        )
        output[f"D_slope_{label}_m2_s"] = slope_D_cm2_s * 1e-4
    finite_slopes = [value for value in slopes if np.isfinite(value) and value > 0]
    output["slope_max_min_ratio"] = (
        max(finite_slopes) / min(finite_slopes) if finite_slopes else float("nan")
    )
    tail = list(state["tail_points"])
    if len(tail) >= 3:
        x = (
            np.asarray([item[0] for item in tail], dtype=float) - float(tail[0][0])
        ) / 60.0
        y = np.asarray([item[1] for item in tail], dtype=float)
        output["rest_tail_abs_slope_mV_min"] = abs(float(np.polyfit(x, y, 1)[0])) * 1000.0
    else:
        output["rest_tail_abs_slope_mV_min"] = float("nan")
    output["qc_early"] = bool(
        output["delta_Es_mV"] >= 2.0
        and output["R2_10_60"] >= 0.99
        and output["slope_max_min_ratio"] <= 1.5
        and output["rest_tail_abs_slope_mV_min"] <= 0.05
    )
    return output


def trim_aggregate(
    rows: list[dict[str, object]], trim_fraction: float, rounding: str
) -> list[dict[str, object]]:
    output: list[dict[str, object]] = []
    for electrode in ("Anode", "Cathode"):
        directions = ("Charge", "Discharge")
        direction_results: dict[str, float] = {}
        for direction in directions:
            cell_values: list[float] = []
            removed: list[int] = []
            for replicate in ("v1", "v2", "v3"):
                group = [
                    item
                    for item in rows
                    if item["electrode"] == electrode
                    and item["direction"] == direction
                    and item["replicate"] == replicate
                ]
                group.sort(key=lambda item: int(item["sequence"]))
                count_float = len(group) * trim_fraction
                count = (
                    math.ceil(count_float)
                    if rounding == "ceil"
                    else math.floor(count_float)
                )
                kept = group[count : len(group) - count if count else None]
                cell_values.append(
                    geometric_mean([float(item["neware_D_m2_s"]) for item in kept])
                )
                removed.append(count)
            direction_results[direction] = geometric_mean(cell_values)
            output.append(
                {
                    "electrode": electrode,
                    "trim_fraction": trim_fraction,
                    "rounding": rounding,
                    "direction": direction,
                    "D_m2_s": direction_results[direction],
                    "removed_each_end_v1_v2_v3": "/".join(map(str, removed)),
                }
            )
        output.append(
            {
                "electrode": electrode,
                "trim_fraction": trim_fraction,
                "rounding": rounding,
                "direction": "Combined",
                "D_m2_s": geometric_mean(list(direction_results.values())),
                "removed_each_end_v1_v2_v3": "-",
            }
        )
    return output


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    columns: list[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def make_figures(rows: list[dict[str, object]], aggregates: list[dict[str, object]]) -> None:
    colors = {"Charge": "#d95f02", "Discharge": "#1b9e77"}
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.6), constrained_layout=True)
    for axis, electrode in zip(axes, ("Anode", "Cathode")):
        for direction in ("Charge", "Discharge"):
            subset = [
                item
                for item in rows
                if item["electrode"] == electrode and item["direction"] == direction
            ]
            x = np.arange(len(subset))
            y = np.asarray([float(item["neware_D_m2_s"]) for item in subset]) / 1e-14
            axis.scatter(x, y, s=10, alpha=0.6, label=direction, color=colors[direction])
        axis.set_yscale("log")
        axis.set_title(f"{electrode}: pulse-wise Neware D")
        axis.set_xlabel("Pulse index (three cells concatenated)")
        axis.set_ylabel(r"$D_s$ [$10^{-14}$ m$^2$/s]")
        axis.grid(True, which="both", alpha=0.25)
        axis.legend()
    fig.savefig(OUT_DIR / "pulsewise_neware_D.png", dpi=180)
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(12, 4.8), constrained_layout=True)
    for axis, electrode in zip(axes, ("Anode", "Cathode")):
        subset = [item for item in aggregates if item["electrode"] == electrode]
        for rounding, linestyle in (("floor", "-"), ("ceil", "--")):
            for direction in ("Charge", "Discharge", "Combined"):
                group = [
                    item
                    for item in subset
                    if item["rounding"] == rounding and item["direction"] == direction
                ]
                group.sort(key=lambda item: float(item["trim_fraction"]))
                axis.plot(
                    [100 * float(item["trim_fraction"]) for item in group],
                    [float(item["D_m2_s"]) / 1e-14 for item in group],
                    linestyle=linestyle,
                    marker="o",
                    ms=3,
                    label=f"{direction}, {rounding}",
                )
        axis.set_title(f"{electrode}: endpoint trimming sensitivity")
        axis.set_xlabel("Trimmed from each end [%]")
        axis.set_ylabel(r"Representative $D_s$ [$10^{-14}$ m$^2$/s]")
        axis.grid(True, alpha=0.25)
        axis.legend(fontsize=8)
    fig.savefig(OUT_DIR / "trim_sensitivity.png", dpi=180)
    plt.close(fig)

    complete = [item for item in rows if item.get("raw_complete")]
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.8), constrained_layout=True)
    for axis, electrode in zip(axes, ("Anode", "Cathode")):
        group = [item for item in complete if item["electrode"] == electrode]
        for window, marker in (("10_60", "o"), ("30_120", "s"), ("30_300", "^")):
            x = np.asarray([float(item["neware_D_m2_s"]) for item in group]) / 1e-14
            y = np.asarray([float(item[f"D_slope_{window}_m2_s"]) for item in group]) / 1e-14
            axis.scatter(x, y, s=13, alpha=0.55, marker=marker, label=window.replace("_", "–") + " s")
        finite = []
        for item in group:
            finite.append(float(item["neware_D_m2_s"]) / 1e-14)
            for window in ("10_60", "30_120", "30_300"):
                finite.append(float(item[f"D_slope_{window}_m2_s"]) / 1e-14)
        finite = [value for value in finite if np.isfinite(value) and value > 0]
        lo, hi = np.percentile(finite, [1, 99]) if finite else (0.1, 100.0)
        axis.plot([lo, hi], [lo, hi], color="black", lw=1, ls=":")
        axis.set_xscale("log")
        axis.set_yscale("log")
        axis.set_title(f"{electrode}: full-pulse vs early-slope D")
        axis.set_xlabel(r"Neware full-pulse $D_s$ [$10^{-14}$ m$^2$/s]")
        axis.set_ylabel(r"Early-slope $D_s$ [$10^{-14}$ m$^2$/s]")
        axis.grid(True, which="both", alpha=0.25)
        axis.legend()
    fig.savefig(OUT_DIR / "full_pulse_vs_early_slope.png", dpi=180)
    plt.close(fig)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    summary_rows: list[dict[str, object]] = []
    reconstructed_rows: list[dict[str, object]] = []
    for electrode_name, replicate, summary_name, raw_name, continuation_name in CASES:
        pulses = read_summary(SUMMARY_DIR / summary_name, electrode_name, replicate)
        summary_rows.extend(pulses)
        raw_paths = [RAW_DIR / raw_name]
        if continuation_name:
            raw_paths.append(RAW_DIR / continuation_name)
        print(f"Reading {electrode_name} {replicate}: {len(pulses)} pulses", flush=True)
        raw = collect_raw_intervals(raw_paths, pulses)
        electrode = ELECTRODES[electrode_name]
        reconstructed_rows.extend(
            reconstruct_pulse(pulse, raw, electrode) for pulse in pulses
        )

    aggregates: list[dict[str, object]] = []
    for fraction in (0.0, 0.05, 0.10, 0.12, 0.15, 0.20):
        for rounding in ("floor", "ceil"):
            aggregates.extend(trim_aggregate(summary_rows, fraction, rounding))

    write_csv(OUT_DIR / "pulse_level_audit.csv", reconstructed_rows)
    write_csv(OUT_DIR / "trim_aggregation_audit.csv", aggregates)

    diagnostics: list[dict[str, object]] = []
    for electrode_name, electrode in ELECTRODES.items():
        group = [
            item
            for item in reconstructed_rows
            if item["electrode"] == electrode_name and item.get("raw_complete")
        ]
        tau_diff = electrode.radius_m**2 / electrode.provisional_D_m2_s
        for direction in ("Charge", "Discharge", "All"):
            subset = group if direction == "All" else [item for item in group if item["direction"] == direction]
            diagnostics.append(
                {
                    "electrode": electrode_name,
                    "direction": direction,
                    "pulse_count": len(subset),
                    "raw_complete_count": len(subset),
                    "neware_reconstruction_abs_error_max_pct": max(
                        [abs(float(item["neware_reconstruction_rel_error_pct"])) for item in subset],
                        default=float("nan"),
                    ),
                    "delta_Es_mV_median": float(np.median([float(item["delta_Es_mV"]) for item in subset])) if subset else float("nan"),
                    "R2_10_60_median": float(np.median([float(item["R2_10_60"]) for item in subset])) if subset else float("nan"),
                    "R2_30_120_median": float(np.median([float(item["R2_30_120"]) for item in subset])) if subset else float("nan"),
                    "R2_30_300_median": float(np.median([float(item["R2_30_300"]) for item in subset])) if subset else float("nan"),
                    "slope_max_min_ratio_median": float(np.median([float(item["slope_max_min_ratio"]) for item in subset])) if subset else float("nan"),
                    "rest_tail_abs_slope_mV_min_median": float(np.median([float(item["rest_tail_abs_slope_mV_min"]) for item in subset])) if subset else float("nan"),
                    "qc_early_pass_count": sum(bool(item["qc_early"]) for item in subset),
                    "D_neware_GM_m2_s": geometric_mean([float(item["neware_D_m2_s"]) for item in subset]),
                    "D_slope_10_60_GM_m2_s": geometric_mean([float(item["D_slope_10_60_m2_s"]) for item in subset]),
                    "D_slope_30_120_GM_m2_s": geometric_mean([float(item["D_slope_30_120_m2_s"]) for item in subset]),
                    "D_slope_30_300_GM_m2_s": geometric_mean([float(item["D_slope_30_300_m2_s"]) for item in subset]),
                    "R2_over_D_s": tau_diff,
                    "Fourier_number_600s": electrode.provisional_D_m2_s * 600.0 / electrode.radius_m**2,
                }
            )
    write_csv(OUT_DIR / "raw_diagnostics_summary.csv", diagnostics)
    make_figures(reconstructed_rows, aggregates)
    print(f"Wrote audit to {OUT_DIR}", flush=True)


if __name__ == "__main__":
    main()
