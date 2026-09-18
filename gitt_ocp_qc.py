"""QC-only preparation pipeline for the 260918 graphite/Li GITT replicates.

This module does not map normalized capacity to absolute stoichiometry, run
Stage 1 with the new candidate, or fit dynamic parameters.
"""

from __future__ import annotations

import bisect
import json
import platform
import warnings
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd


warnings.filterwarnings("ignore", message="Workbook contains no default style")

RAW_FILENAMES = {
    "v2": "260918 Anode GITT v2 1-6.xlsx",
    "v3": "260918 Anode GITT v3 1-7.xlsx",
}
RESULT_DIRNAME = "260918_gitt_ocp_qc"
ANODE_OCP_CASE = "legacy"


@dataclass(frozen=True)
class QCPaths:
    root: Path
    raw_dir: Path
    legacy_book: Path
    results: Path


def get_paths(root: Path | str | None = None) -> QCPaths:
    root_path = Path(root or Path.cwd()).resolve()
    return QCPaths(
        root=root_path,
        raw_dir=root_path / "data" / "raw" / "260918_Anode_GITT",
        legacy_book=root_path / "260902 Enertech셀 데이터 모음.xlsx",
        results=root_path / "results" / RESULT_DIRNAME,
    )


def _duration_seconds(value: Any) -> float:
    if isinstance(value, timedelta):
        return value.total_seconds()
    if isinstance(value, (int, float)):
        return float(value) * 86400.0
    pieces = str(value).split(":")
    if len(pieces) != 3:
        raise ValueError(f"Unsupported duration: {value!r}")
    hours, minutes, seconds = (float(piece) for piece in pieces)
    return hours * 3600.0 + minutes * 60.0 + seconds


def _parse_datetime(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value
    return datetime.fromisoformat(str(value))


def _stream_rows(workbook: openpyxl.Workbook, sheet_name: str):
    sheet = workbook[sheet_name]
    sheet.reset_dimensions()
    yield from sheet.iter_rows(values_only=True)


def _read_step_table(workbook: openpyxl.Workbook, replicate: str) -> pd.DataFrame:
    rows = list(_stream_rows(workbook, "step"))
    frame = pd.DataFrame(rows[1:], columns=rows[0])
    frame["replicate"] = replicate
    frame["Step Number"] = pd.to_numeric(frame["Step Number"], errors="raise").astype(int)
    frame["Cycle Index"] = pd.to_numeric(frame["Cycle Index"], errors="raise").astype(int)
    frame["Capacity(Ah)"] = pd.to_numeric(frame["Capacity(Ah)"], errors="coerce")
    frame["Oneset Date"] = pd.to_datetime(frame["Oneset Date"])
    frame["End Date"] = pd.to_datetime(frame["End Date"])
    frame["step_duration_s"] = frame["Step Time"].map(_duration_seconds)
    frame["measured_duration_s"] = (
        frame["End Date"] - frame["Oneset Date"]
    ).dt.total_seconds()
    return frame


def _find_gitt_start(steps: pd.DataFrame) -> int:
    candidates = steps[
        steps["Cycle Index"].eq(2)
        & steps["Step Type"].eq("Rest")
        & steps["step_duration_s"].between(7100, 7300)
    ]
    if candidates.empty:
        raise RuntimeError("No two-hour GITT rest found")
    return int(candidates.index[0])


def _collect_target_records(
    workbook: openpyxl.Workbook,
    steps: pd.DataFrame,
    target_step_numbers: set[int],
) -> dict[int, dict[str, list[float]]]:
    starts = [value.to_pydatetime() for value in steps["Oneset Date"]]
    ends = [value.to_pydatetime() for value in steps["End Date"]]
    numbers = steps["Step Number"].astype(int).tolist()
    output = {
        number: {"elapsed_min": [], "voltage_V": [], "current_A": []}
        for number in target_step_numbers
    }
    rows = _stream_rows(workbook, "record")
    header = next(rows)
    index = {name: position for position, name in enumerate(header)}
    for values in rows:
        raw_date = values[index["Date"]]
        raw_voltage = values[index["Voltage(V)"]]
        raw_current = values[index["Current(A)"]]
        if raw_date is None or raw_voltage is None:
            continue
        stamp = _parse_datetime(raw_date)
        position = bisect.bisect_right(starts, stamp) - 1
        if position < 0 or position >= len(numbers):
            continue
        if stamp > ends[position] + timedelta(seconds=2):
            continue
        step_number = numbers[position]
        if step_number not in target_step_numbers:
            continue
        output[step_number]["elapsed_min"].append(
            (stamp - starts[position]).total_seconds() / 60.0
        )
        output[step_number]["voltage_V"].append(float(raw_voltage))
        output[step_number]["current_A"].append(float(raw_current or 0.0))
    return output


def _interp_at(elapsed_min: np.ndarray, values: np.ndarray, minute: float) -> float:
    order = np.argsort(elapsed_min)
    return float(np.interp(minute, elapsed_min[order], values[order]))


def _rest_metrics(samples: dict[str, list[float]]) -> dict[str, float]:
    t = np.asarray(samples["elapsed_min"], dtype=float)
    voltage = np.asarray(samples["voltage_V"], dtype=float)
    valid = np.isfinite(t) & np.isfinite(voltage)
    t = t[valid]
    voltage = voltage[valid]
    if len(t) < 3:
        return {
            "rest_start_voltage_V": np.nan,
            "rest_voltage_at_60_min_V": np.nan,
            "rest_end_voltage_at_120_min_V": np.nan,
            "delta_V_2h_minus_1h_mV": np.nan,
            "last_10_min_slope_mV_per_min": np.nan,
            "last_10_min_abs_dVdt_mV_per_min": np.nan,
            "rest_record_count": int(len(t)),
        }
    order = np.argsort(t)
    t = t[order]
    voltage = voltage[order]
    v_1h = _interp_at(t, voltage, 60.0)
    v_2h = _interp_at(t, voltage, min(120.0, float(t[-1])))
    tail = t >= max(float(t[-1]) - 10.0, float(t[0]))
    slope_v_per_min = float(np.polyfit(t[tail], voltage[tail], 1)[0]) if tail.sum() >= 3 else np.nan
    return {
        "rest_start_voltage_V": float(voltage[0]),
        "rest_voltage_at_60_min_V": v_1h,
        "rest_end_voltage_at_120_min_V": v_2h,
        "delta_V_2h_minus_1h_mV": (v_2h - v_1h) * 1000.0,
        "last_10_min_slope_mV_per_min": slope_v_per_min * 1000.0,
        "last_10_min_abs_dVdt_mV_per_min": abs(slope_v_per_min) * 1000.0,
        "rest_record_count": int(len(t)),
    }


def parse_replicate(path: Path, replicate: str) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    required = {"step", "record"}
    missing = required.difference(workbook.sheetnames)
    if missing:
        raise RuntimeError(f"{path.name}: missing sheets {sorted(missing)}")
    steps = _read_step_table(workbook, replicate)
    gitt_start = _find_gitt_start(steps)
    gitt_steps = steps.loc[gitt_start:].copy().reset_index(drop=True)
    target_numbers = set(gitt_steps["Step Number"].astype(int))
    records = _collect_target_records(workbook, steps, target_numbers)

    pulse_rows: list[dict[str, Any]] = []
    rest_rows: list[dict[str, Any]] = []
    counters = {"lithiation": 0, "delithiation": 0}
    cumulative = {"lithiation": 0.0, "delithiation": 0.0}
    last_direction = "lithiation"
    last_pulse_no = 0
    last_cumulative = 0.0

    for position, row in gitt_steps.iterrows():
        step_type = str(row["Step Type"])
        step_number = int(row["Step Number"])
        if step_type in {"CC DChg", "CC Chg"}:
            direction = "lithiation" if step_type == "CC DChg" else "delithiation"
            counters[direction] += 1
            capacity_mAh = float(row["Capacity(Ah)"]) * 1000.0
            cumulative[direction] += capacity_mAh
            samples = records.get(step_number, {})
            current = np.asarray(samples.get("current_A", []), dtype=float)
            pulse_current_mA = float(np.nanmedian(current) * 1000.0) if len(current) else np.nan
            next_row = gitt_steps.iloc[position + 1] if position + 1 < len(gitt_steps) else None
            rest_data: dict[str, float] = {}
            rest_step_number = np.nan
            rest_start = pd.NaT
            rest_end = pd.NaT
            rest_duration_min = np.nan
            if (
                next_row is not None
                and str(next_row["Step Type"]) == "Rest"
                and 7100 <= float(next_row["step_duration_s"]) <= 7300
            ):
                rest_step_number = int(next_row["Step Number"])
                rest_start = next_row["Oneset Date"]
                rest_end = next_row["End Date"]
                rest_duration_min = float(next_row["measured_duration_s"]) / 60.0
                rest_data = _rest_metrics(records.get(rest_step_number, {}))
            pulse_rows.append(
                {
                    "cell": replicate,
                    "direction": direction,
                    "pulse_number": counters[direction],
                    "pulse_step_number": step_number,
                    "pulse_start_time": row["Oneset Date"],
                    "pulse_end_time": row["End Date"],
                    "pulse_duration_min": float(row["measured_duration_s"]) / 60.0,
                    "pulse_current_mA_signed": pulse_current_mA,
                    "pulse_current_mA_abs": abs(pulse_current_mA),
                    "pulse_capacity_increment_mAh": capacity_mAh,
                    "cumulative_branch_capacity_mAh": cumulative[direction],
                    "pulse_start_voltage_V": float(row["Oneset Volt.(V)"]),
                    "pulse_end_voltage_V": float(row["End Voltage(V)"]),
                    "rest_step_number": rest_step_number,
                    "rest_start_time": rest_start,
                    "rest_end_time": rest_end,
                    "rest_duration_min": rest_duration_min,
                    **rest_data,
                }
            )
            last_direction = direction
            last_pulse_no = counters[direction]
            last_cumulative = cumulative[direction]
        elif step_type == "Rest" and 7100 <= float(row["step_duration_s"]) <= 7300:
            rest_rows.append(
                {
                    "cell": replicate,
                    "direction": last_direction,
                    "pulse_number": last_pulse_no,
                    "cumulative_branch_capacity_mAh": last_cumulative,
                    "rest_step_number": step_number,
                    "rest_start_time": row["Oneset Date"],
                    "rest_end_time": row["End Date"],
                    "rest_duration_min": float(row["measured_duration_s"]) / 60.0,
                    "rest_context": "pre_gitt_initial" if last_pulse_no == 0 else "after_pulse",
                    **_rest_metrics(records.get(step_number, {})),
                }
            )

    pulse_frame = pd.DataFrame(pulse_rows)
    for direction in ("lithiation", "delithiation"):
        mask = pulse_frame["direction"].eq(direction)
        branch_capacity = pulse_frame.loc[mask, "cumulative_branch_capacity_mAh"].max()
        pulse_frame.loc[mask, "branch_usable_capacity_mAh"] = branch_capacity
        pulse_frame.loc[mask, "q_norm_branch"] = (
            pulse_frame.loc[mask, "cumulative_branch_capacity_mAh"] / branch_capacity
        )
    rest_frame = pd.DataFrame(rest_rows)
    for direction in ("lithiation", "delithiation"):
        mask = rest_frame["direction"].eq(direction)
        branch_capacity = pulse_frame.loc[
            pulse_frame["direction"].eq(direction), "branch_usable_capacity_mAh"
        ].max()
        rest_frame.loc[mask, "q_norm_branch"] = (
            rest_frame.loc[mask, "cumulative_branch_capacity_mAh"] / branch_capacity
        )

    return {
        "path": path,
        "steps": steps,
        "gitt_steps": gitt_steps,
        "pulse_rest": pulse_frame,
        "rests": rest_frame,
        "parser_status": {
            "cell": replicate,
            "source_file": path.name,
            "sheet_names": ", ".join(workbook.sheetnames),
            "step_count": int(len(steps)),
            "gitt_step_count": int(len(gitt_steps)),
            "pulse_count": int(len(pulse_frame)),
            "two_hour_rest_count": int(len(rest_frame)),
            "record_mapping_complete": bool(
                pulse_frame["rest_record_count"].fillna(0).gt(0).all()
            ),
            "status": "OK",
        },
    }


def _add_region_rows(frame: pd.DataFrame, value_col: str) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for (cell, direction), group in frame.groupby(["cell", "direction"], sort=True):
        regions = {
            "full": np.ones(len(group), dtype=bool),
            "central_10_90": group["q_norm_branch"].between(0.10, 0.90).to_numpy(),
            "endpoints": (~group["q_norm_branch"].between(0.10, 0.90)).to_numpy(),
        }
        for region, mask in regions.items():
            values = pd.to_numeric(group.loc[mask, value_col], errors="coerce").dropna().to_numpy()
            if len(values) == 0:
                continue
            absolute = np.abs(values)
            rows.append(
                {
                    "cell": cell,
                    "direction": direction,
                    "region": region,
                    "n": int(len(values)),
                    "mean": float(np.mean(values)),
                    "median": float(np.median(values)),
                    "mean_abs": float(np.mean(absolute)),
                    "median_abs": float(np.median(absolute)),
                    "p90_abs": float(np.percentile(absolute, 90)),
                    "max_abs": float(np.max(absolute)),
                }
            )
    return pd.DataFrame(rows)


def _capacity_qc(parsed: dict[str, dict[str, Any]]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for cell, item in parsed.items():
        steps = item["steps"]
        gitt_start_step = int(item["gitt_steps"].iloc[0]["Step Number"])
        initial = steps[steps["Step Number"] < gitt_start_step]
        lith = initial[initial["Step Type"].eq("CC DChg")]["Capacity(Ah)"].to_numpy() * 1000.0
        delith = initial[initial["Step Type"].eq("CC Chg")]["Capacity(Ah)"].to_numpy() * 1000.0
        pulse = item["pulse_rest"]
        gitt_lith = float(
            pulse.loc[pulse["direction"].eq("lithiation"), "pulse_capacity_increment_mAh"].sum()
        )
        gitt_delith = float(
            pulse.loc[pulse["direction"].eq("delithiation"), "pulse_capacity_increment_mAh"].sum()
        )
        rows.append(
            {
                "cell": cell,
                "initial_lithiation_cycle1_mAh": float(lith[0]) if len(lith) > 0 else np.nan,
                "initial_delithiation_cycle1_mAh": float(delith[0]) if len(delith) > 0 else np.nan,
                "initial_lithiation_cycle2_mAh": float(lith[1]) if len(lith) > 1 else np.nan,
                "initial_delithiation_cycle2_mAh": float(delith[1]) if len(delith) > 1 else np.nan,
                "gitt_lithiation_capacity_mAh": gitt_lith,
                "gitt_delithiation_capacity_mAh": gitt_delith,
                "gitt_reversible_capacity_mean_mAh": 0.5 * (gitt_lith + gitt_delith),
                "gitt_lith_minus_delith_mAh": gitt_lith - gitt_delith,
                "gitt_delith_over_lith_ratio": gitt_delith / gitt_lith,
                "full_cell_window_reference_low_mAh": 2.20,
                "full_cell_window_reference_high_mAh": 2.25,
                "automatic_capacity_correction_applied": False,
            }
        )
    output = pd.DataFrame(rows)
    for column in ("gitt_lithiation_capacity_mAh", "gitt_delithiation_capacity_mAh"):
        values = output.set_index("cell")[column]
        difference = abs(float(values["v2"] - values["v3"])) / float(values.mean()) * 100.0
        output[f"v2_v3_{column}_difference_pct"] = difference
    return output


def _clean_curve(x: np.ndarray, y: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    frame = pd.DataFrame({"x": x, "y": y}).replace([np.inf, -np.inf], np.nan).dropna()
    grouped = frame.groupby("x", as_index=False)["y"].mean().sort_values("x")
    return grouped["x"].to_numpy(float), grouped["y"].to_numpy(float)


def _difference_metrics(a: np.ndarray, b: np.ndarray) -> dict[str, float]:
    difference_mV = (a - b) * 1000.0
    absolute = np.abs(difference_mV)
    return {
        "MAE_mV": float(np.mean(absolute)),
        "RMSE_mV": float(np.sqrt(np.mean(difference_mV**2))),
        "median_absolute_difference_mV": float(np.median(absolute)),
        "P90_absolute_difference_mV": float(np.percentile(absolute, 90)),
        "max_absolute_difference_mV": float(np.max(absolute)),
    }


def _normalized_new_curves(pulse_rest: pd.DataFrame):
    grid = np.linspace(0.0, 1.0, 1001)
    curves: dict[tuple[str, str], np.ndarray] = {}
    metrics: list[dict[str, Any]] = []
    for direction in ("lithiation", "delithiation"):
        for cell in ("v2", "v3"):
            selected = pulse_rest[
                pulse_rest["cell"].eq(cell) & pulse_rest["direction"].eq(direction)
            ]
            x, y = _clean_curve(
                selected["q_norm_branch"].to_numpy(float),
                selected["rest_end_voltage_at_120_min_V"].to_numpy(float),
            )
            curves[(cell, direction)] = np.interp(grid, x, y)
        for region, mask in (
            ("full_overlap", np.ones_like(grid, dtype=bool)),
            ("central_10_90", (grid >= 0.10) & (grid <= 0.90)),
        ):
            metrics.append(
                {
                    "comparison": "v2_minus_v3",
                    "direction": direction,
                    "region": region,
                    **_difference_metrics(
                        curves[("v2", direction)][mask],
                        curves[("v3", direction)][mask],
                    ),
                }
            )
    lith_mean = 0.5 * (curves[("v2", "lithiation")] + curves[("v3", "lithiation")])
    delith_progress_mean = 0.5 * (
        curves[("v2", "delithiation")] + curves[("v3", "delithiation")]
    )
    position = grid
    delith_position_curve = delith_progress_mean[::-1]
    representative = 0.5 * (lith_mean + delith_position_curve)
    candidate = pd.DataFrame(
        {
            "Q_NORM_NEW": position,
            "UN_NEW_LITH_MEAN_V": lith_mean,
            "UN_NEW_DELITH_MEAN_V": delith_position_curve,
            "UN_NEW_NORM_CANDIDATE_V": representative,
            "axis_interpretation": "normalized_position_not_absolute_stoichiometry",
        }
    )
    return grid, curves, pd.DataFrame(metrics), candidate


def _legacy_normalized_curve(legacy_book: Path, grid: np.ndarray) -> tuple[np.ndarray, dict]:
    frame = pd.read_excel(legacy_book, sheet_name="ANODE_LOWRATE", header=8)
    curves = []
    detail = {}
    for source in ("v1", "v2"):
        for direction in ("discharge", "charge"):
            selected = frame[
                frame["source_version"].astype(str).eq(source)
                & frame["direction"].astype(str).eq(direction)
            ]
            capacity = pd.to_numeric(selected["Q_Ah"], errors="coerce").to_numpy(float)
            voltage = pd.to_numeric(selected["V_V"], errors="coerce").to_numpy(float)
            q_norm = capacity / np.nanmax(capacity)
            position = q_norm if direction == "discharge" else 1.0 - q_norm
            x, y = _clean_curve(position, voltage)
            interpolated = np.interp(grid, x, y)
            detail[(source, direction)] = interpolated
            curves.append(interpolated)
    return np.mean(np.vstack(curves), axis=0), detail


def _legacy_comparison(
    grid: np.ndarray,
    legacy_curve: np.ndarray,
    candidate: pd.DataFrame,
) -> pd.DataFrame:
    new_curve = candidate["UN_NEW_NORM_CANDIDATE_V"].to_numpy(float)
    rows = []
    for region, mask in (
        ("full_overlap", np.ones_like(grid, dtype=bool)),
        ("central_10_90", (grid >= 0.10) & (grid <= 0.90)),
    ):
        rows.append(
            {
                "comparison": "legacy_minus_new_normalized_candidate",
                "region": region,
                "axis_interpretation": "normalized_shape_only_not_absolute_stoichiometry",
                **_difference_metrics(legacy_curve[mask], new_curve[mask]),
            }
        )
    return pd.DataFrame(rows)


def stage1_comparison_skeleton() -> pd.DataFrame:
    columns = [
        "case", "anode_source", "cathode_source", "status", "blocking_reason",
        "x0", "x100", "y0", "y100", "DX", "DY",
        "endpoint_0_error_mV", "endpoint_100_error_mV",
        "qOCV_RMSE_full_mV", "qOCV_MAE_full_mV",
        "qOCV_RMSE_2_98_mV", "qOCV_MAE_2_98_mV",
    ]
    rows = [
        ["A", "legacy", "legacy_current", "READY_EXISTING_METHOD", "", *([np.nan] * 12)],
        ["B", "new_2h_candidate", "legacy_current", "NOT_READY", "absolute anode stoichiometry mapping not established", *([np.nan] * 12)],
        ["C", "legacy", "new_2h_candidate", "NOT_READY", "final cathode 2 h GITT data unavailable", *([np.nan] * 12)],
        ["D", "new_2h_candidate", "new_2h_candidate", "NOT_READY", "absolute anode mapping and final cathode data unavailable", *([np.nan] * 12)],
    ]
    return pd.DataFrame(rows, columns=columns)


def select_anode_ocp_source(case: str, absolute_mapping: Any = None) -> dict[str, str]:
    if case == "legacy":
        return {"case": case, "status": "READY_EXISTING_METHOD"}
    if case == "new_2h_candidate" and absolute_mapping is None:
        return {
            "case": case,
            "status": "NOT_READY",
            "reason": "Q_NORM_NEW is not absolute stoichiometry; Stage 1 injection blocked",
        }
    if case != "new_2h_candidate":
        raise ValueError(f"Unsupported ANODE_OCP_CASE: {case}")
    return {"case": case, "status": "READY_WITH_USER_SUPPLIED_ABSOLUTE_MAPPING"}


def _save_figures(
    results: Path,
    pulse_rest: pd.DataFrame,
    rests: pd.DataFrame,
    grid: np.ndarray,
    curves: dict,
    candidate: pd.DataFrame,
    legacy_curve: np.ndarray,
):
    colors = {"v2": "#1f77b4", "v3": "#d62728"}
    markers = {"lithiation": "o", "delithiation": "s"}

    fig, axes = plt.subplots(1, 2, figsize=(11, 4.2), sharey=True)
    for axis, direction in zip(axes, ("lithiation", "delithiation")):
        for cell in ("v2", "v3"):
            selected = pulse_rest[pulse_rest["cell"].eq(cell) & pulse_rest["direction"].eq(direction)]
            axis.plot(selected["q_norm_branch"], selected["delta_V_2h_minus_1h_mV"], marker=markers[direction], ms=3, lw=1.0, label=cell, color=colors[cell])
        axis.axhline(0.0, color="black", lw=0.8)
        axis.set(title=direction.capitalize(), xlabel="Normalized branch capacity")
    axes[0].set_ylabel("V(2 h) - V(1 h) [mV]")
    axes[0].legend()
    fig.tight_layout()
    fig.savefig(results / "anode_rest_1h_vs_2h.png", dpi=180)
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(11, 4.2), sharey=True)
    for axis, direction in zip(axes, ("lithiation", "delithiation")):
        for cell in ("v2", "v3"):
            selected = rests[
                rests["cell"].eq(cell)
                & rests["direction"].eq(direction)
                & rests["rest_context"].eq("after_pulse")
            ]
            displayed_slope = np.maximum(
                selected["last_10_min_abs_dVdt_mV_per_min"].to_numpy(float),
                1e-4,
            )
            axis.plot(selected["q_norm_branch"], displayed_slope, marker=markers[direction], ms=3, lw=1.0, label=cell, color=colors[cell])
        axis.set(title=direction.capitalize(), xlabel="Normalized branch capacity")
        axis.set_yscale("log")
        axis.set_ylim(1e-4, 1.0)
    axes[0].set_ylabel("Last-10-min |dV/dt| [mV/min] (display floor 1e-4)")
    axes[0].legend()
    fig.tight_layout()
    fig.savefig(results / "anode_rest_equilibrium.png", dpi=180)
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(11, 4.2), sharey=True)
    for axis, direction in zip(axes, ("lithiation", "delithiation")):
        for cell in ("v2", "v3"):
            axis.plot(grid, curves[(cell, direction)], lw=1.6, label=cell, color=colors[cell])
        axis.set(title=direction.capitalize(), xlabel="Normalized branch capacity", ylabel="Rest-end voltage [V]")
        axis.legend()
    fig.tight_layout()
    fig.savefig(results / "anode_v2_v3_normalized_reproducibility.png", dpi=180)
    plt.close(fig)

    fig, axis = plt.subplots(figsize=(7.5, 5.0))
    axis.plot(grid, legacy_curve, color="black", lw=2.0, label="Legacy normalized mean")
    axis.plot(grid, candidate["UN_NEW_LITH_MEAN_V"], lw=1.2, label="New lithiation mean")
    axis.plot(grid, candidate["UN_NEW_DELITH_MEAN_V"], lw=1.2, label="New delithiation mean")
    axis.plot(grid, candidate["UN_NEW_NORM_CANDIDATE_V"], lw=2.0, label="New representative candidate")
    axis.set(xlabel="Normalized position (not absolute stoichiometry)", ylabel="Anode potential [V vs. Li/Li+]", title="Legacy vs new anode OCP shape")
    axis.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(results / "legacy_vs_new_anode_ocp_normalized.png", dpi=180)
    plt.close(fig)


def run_qc(root: Path | str | None = None) -> dict[str, Any]:
    paths = get_paths(root)
    paths.results.mkdir(parents=True, exist_ok=True)
    parsed = {
        cell: parse_replicate(paths.raw_dir / filename, cell)
        for cell, filename in RAW_FILENAMES.items()
    }
    pulse_rest = pd.concat([item["pulse_rest"] for item in parsed.values()], ignore_index=True)
    rests = pd.concat([item["rests"] for item in parsed.values()], ignore_index=True)
    parser_status = pd.DataFrame([item["parser_status"] for item in parsed.values()])
    capacity_qc = _capacity_qc(parsed)
    relaxation = pulse_rest[
        [
            "cell", "direction", "pulse_number", "q_norm_branch",
            "rest_voltage_at_60_min_V", "rest_end_voltage_at_120_min_V",
            "delta_V_2h_minus_1h_mV", "rest_duration_min",
        ]
    ].copy()
    relaxation_metrics = _add_region_rows(pulse_rest, "delta_V_2h_minus_1h_mV")
    equilibrium = rests.copy()
    equilibrium_metrics = _add_region_rows(
        rests[rests["rest_context"].eq("after_pulse")],
        "last_10_min_slope_mV_per_min",
    )
    grid, curves, reproducibility, candidate = _normalized_new_curves(pulse_rest)
    legacy_curve, legacy_detail = _legacy_normalized_curve(paths.legacy_book, grid)
    legacy_metrics = _legacy_comparison(grid, legacy_curve, candidate)
    stage1_skeleton = stage1_comparison_skeleton()
    cathode_placeholder = pd.DataFrame(
        [{
            "component": "cathode_2h_gitt",
            "status": "WAITING_FOR_FINAL_RAW_DATA",
            "planned_pipeline": "pulse/rest parser; 1h-vs-2h; last-10-min dV/dt; replicate QC; normalized OCP; legacy comparison",
            "final_ocp_used": False,
        }]
    )

    csv_outputs = {
        "anode_pulse_rest_summary.csv": pulse_rest,
        "anode_1h_vs_2h_relaxation.csv": relaxation,
        "anode_1h_vs_2h_relaxation_metrics.csv": relaxation_metrics,
        "anode_rest_equilibrium.csv": equilibrium,
        "anode_rest_equilibrium_metrics.csv": equilibrium_metrics,
        "anode_capacity_qc.csv": capacity_qc,
        "anode_v2_v3_reproducibility.csv": reproducibility,
        "anode_normalized_ocp_candidate.csv": candidate,
        "legacy_vs_new_anode_ocp_metrics.csv": legacy_metrics,
        "stage1_comparison_skeleton.csv": stage1_skeleton,
        "cathode_2h_placeholder.csv": cathode_placeholder,
        "parser_status.csv": parser_status,
    }
    for filename, frame in csv_outputs.items():
        frame.to_csv(paths.results / filename, index=False, encoding="utf-8-sig")

    _save_figures(paths.results, pulse_rest, rests, grid, curves, candidate, legacy_curve)
    manifest = {
        "python_version": platform.python_version(),
        "raw_files": RAW_FILENAMES,
        "anode_ocp_case_default": ANODE_OCP_CASE,
        "new_candidate_axis": "Q_NORM_NEW; normalized position only",
        "new_candidate_stage1_status": select_anode_ocp_source("new_2h_candidate"),
        "absolute_stoichiometry_anchoring_performed": False,
        "final_ocp_selection_performed": False,
        "stage2_to_stage4_fitting_performed": False,
    }
    (paths.results / "qc_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return {
        "paths": paths,
        "parsed": parsed,
        "parser_status": parser_status,
        "pulse_rest": pulse_rest,
        "relaxation": relaxation,
        "relaxation_metrics": relaxation_metrics,
        "equilibrium": equilibrium,
        "equilibrium_metrics": equilibrium_metrics,
        "capacity_qc": capacity_qc,
        "grid": grid,
        "curves": curves,
        "reproducibility": reproducibility,
        "candidate": candidate,
        "legacy_curve": legacy_curve,
        "legacy_detail": legacy_detail,
        "legacy_metrics": legacy_metrics,
        "stage1_skeleton": stage1_skeleton,
        "cathode_placeholder": cathode_placeholder,
        "manifest": manifest,
    }


if __name__ == "__main__":
    run_qc()
